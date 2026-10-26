# -*- coding: utf-8 -*-
# x51: V11 —— ①堆叠柱(内/外双色+段上数值) ②D1页新增"纯提PR未合入"柱
import json, os, re
from collections import Counter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn as _qn

B=os.path.join(os.path.dirname(os.path.abspath(__file__)),"data")+os.sep
DEL="/Users/tian/Desktop/开源/昇腾社区26年1月-9月贡献者分析/最新交付/"
OUT=DEL+"9社区流失开发者分析_V11.pptx"
COMMS=["vllm","sgl","VeRL","triton","MindIE","MindSDK","MindStudio","MindCluster","MindSpeed"]
X=json.load(open(B+"x40_data.json")); _=X.pop("CATS")
TM=json.load(open(B+"task_meta.json")); AL=json.load(open(B+"active_latest.json"))
AD=json.load(open(B+"active_detail.json")); TI=json.load(open(B+"titles_1021.json")); CM=json.load(open(B+"compmap.json"))
D0={c:[0]*7+[AL[c]["aug"]["D0"]["all"],AL[c]["sep"]["D0"]["all"]] for c in COMMS}

# ---- 分类规则(与x40一致) + 逐人内/外 & 纯提PR ----
BUG_K=["[bug","bugfix","bug-fix","修复","缺陷","报错","崩溃","crash","error","失败","故障","异常","oom","segfault","挂死","hang","乱码","精度","broken","corrupt","wrong","incorrect","fail","死循环","重复输出","garbled","timeout","fix","resolve","hotfix","patch"]
FEAT_K=["[feature","[rfc","[new model","新增","support","添加特性","引入","feature-request","需求反馈","roadmap","请求"]
DOC_K=["[doc","documentation","文档","readme","tutorial"]; PERF_K=["[performance","[perf","性能","optimize","优化","regression","劣化","throughput","ttft","tpot"]
CI_K=["[ci","[test","测试","nightly","daily-ci","test case","e2e","ci门禁"]; RF_K=["[refactor","重构","refactor","retire","sunset"]
IN_K=["[installation","安装","部署","installation"]; TASK_K=["淬火行动","社区任务","任务完成","测试报告","昇腾开源社区任务","向我开炮"]
CATS=["修复Bug","报告Bug","纯提PR未合入","评论型","社区任务","新特性","文档","性能优化","CI/测试","重构","安装部署","其他"]
def _tag(t):
    m=re.match(r'^\[([^\]]+)\]',t); g=m.group(1).lower() if m else ""
    return g.startswith("bug") or "bug-report" in g
def classify_pr(title):
    if not title: return set()
    t=title.lower(); c=set()
    if any(k in t for k in BUG_K) or _tag(title): c.add("修复Bug")
    if any(k in t for k in FEAT_K) or _tag(title): c.add("新特性")
    if any(k in t for k in TASK_K): c.add("社区任务")
    if any(k in t for k in DOC_K): c.add("文档")
    if any(k in t for k in PERF_K): c.add("性能优化")
    if any(k in t for k in CI_K): c.add("CI/测试")
    if any(k in t for k in RF_K): c.add("重构")
    if any(k in t for k in IN_K): c.add("安装部署")
    if not c: c.add("其他")
    return c
def classify_iss(title):
    if not title: return set()
    t=title.lower(); c=set()
    if any(k in t for k in TASK_K): c.add("社区任务")
    elif any(k in t for k in BUG_K) or _tag(title): c.add("报告Bug")
    if any(k in t for k in FEAT_K): c.add("新特性")
    if any(k in t for k in DOC_K): c.add("文档")
    if any(k in t for k in PERF_K): c.add("性能优化")
    if any(k in t for k in CI_K): c.add("CI/测试")
    if any(k in t for k in IN_K): c.add("安装部署")
    if not c: c.add("其他")
    return c
def build(c,lv):
    lost=set(X[c][lv]['lost']); isin=lambda u:"华为" in (CM.get(c,{}).get(u) or "")
    ev={}
    for u,ty,_ in AD[c][f"8-{lv}"]: ev.setdefault(u,set()).add(ty)
    ci=Counter(); ce=Counter()
    for u in sorted(lost):
        ct=set()
        for ti in TI[c]["aug"]["pr"].get(u,[]): ct|=classify_pr(ti)
        for ti in TI[c]["aug"]["iss"].get(u,[]): ct|=classify_iss(ti)
        types=ev.get(u,set())
        if not types and (TI[c]["aug"]["pr"].get(u) or TI[c]["aug"]["iss"].get(u)): types={"pr","issue"}
        if lv=="D1" and types=={"pr"}: ct.add("纯提PR未合入")
        if "pr" in types or "issue" in types:
            pass
        if not ct and not ("pr" in types or "issue" in types): ct={"评论型"} if "comment" in types else {"其他"}
        if "comment" in types: ct.add("评论型")
        tgt=ci if isin(u) else ce
        for k in ct: tgt[k]+=1
    return ci,ce
GRP={}
for c in COMMS:
    GRP[c]={lv:build(c,lv) for lv in ("D1","D2")}

# ---- 渲染 ----
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
def tb(s,l,tp,w,h,text,size=14,bold=False,color=RGBColor(0x33,0x33,0x33),align=PP_ALIGN.LEFT,paras=None):
    box=s.shapes.add_textbox(l,tp,w,h); tf=box.text_frame; tf.word_wrap=True
    for i,seg in enumerate(paras if paras else [text]):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=seg; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box
def pct(a,b): return f"{a*100.0/b:.1f}%" if b else "—"
C_IN=RGBColor(0x8F,0x2D,0x2D); C_EXT=RGBColor(0x2E,0x74,0xB5)
for c in COMMS:
    for lv in ("D1","D2"):
        o=X[c][lv]; ci,ce=GRP[c][lv]
        s=prs.slides.add_slide(blank)
        tb(s,Inches(0.4),Inches(0.2),Inches(12.5),Inches(0.6),f"{c} — {lv} 活跃开发者流失分析（8月→9月）",28,True,RGBColor(0x1F,0x38,0x64))
        ex8=o["p8"]-o["p8_in"]; ex9=o["p9"]-o["p9_in"]; dn=o["p9"]-o["p8"]
        red=[f"{lv} 活跃开发者总数：8月 {o['p8']} 人（内部 {o['p8_in']} 人、外部 {ex8} 人）→ 9月 {o['p9']} 人（内部 {o['p9_in']} 人、外部 {ex9} 人）",
             f"{lv} 活跃开发者9月流失：-{o['lost_n']} 人（内部 -{o['lost_in']} 人、外部 -{o['lost_ext']} 人）　|　净变化 {dn:+d} 人（{'升入D2 '+str(o.get('up',0))+'人、自D2降级流入 '+str(o.get('down',0))+'人' if lv=='D1' else '自D1升入 '+str(o.get('in_up',0))+'人、降级流出至D1 '+str(o.get('out_down',0))+'人'}，层间净{'流出' if (dn-(o['new_n']-o['lost_n']))<0 else '流入'} {abs(dn-(o['new_n']-o['lost_n']))} 人）",
             f"{lv} 活跃开发者9月新增：{o['new_n']} 人（内部 {o['new_in']} 人、外部 {o['new_ext']} 人）"]
        tb(s,Inches(0.4),Inches(0.88),Inches(12.6),Inches(2.2),None,26,True,RGBColor(0xC0,0x00,0x00),paras=red)
        tb(s,Inches(0.4),Inches(3.05),Inches(12.6),Inches(0.4),
           "口径：当月活跃池逐人点名（D1＝PR/Issue/评论任一，D2＝有合入PR）。留存 %d 人（%s）；账目：8月池 %d ＝ 留存%d＋流失%d＋%s；9月池 %d ＝ 留存%d＋新增%d＋%s"%(
             o['stay'],pct(o['stay'],o['p8']),o['p8'],o['stay'],o['lost_n'],('升入D2 '+str(o.get('up',0))) if lv=='D1' else ('降级流出 '+str(o.get('out_down',0))),
             o['p9'],o['stay'],o['new_n'],('自D2降级 '+str(o.get('down',0))) if lv=='D1' else ('自D1升入 '+str(o.get('in_up',0)))),
           11,False,RGBColor(0x88,0x88,0x88))
        cats=[k for k in CATS if (lv=="D1" or k!="纯提PR未合入") and (ci.get(k,0)+ce.get(k,0))>0]
        if cats:
            rev=list(reversed(cats))
            cd=CategoryChartData(); cd.categories=rev
            cd.add_series("外部",  [ce.get(k,0) for k in rev])
            cd.add_series("内部",  [ci.get(k,0) for k in rev])
            gf=s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED,Inches(0.4),Inches(3.5),Inches(7.4),Inches(3.35),cd)
            ch=gf.chart; ch.has_title=True
            ch.chart_title.text_frame.text=f"9月流失开发者 · 8月贡献活动分布（流失{o['lost_n']}人，柱内=内部■{C_IN}／外部■蓝，数字为人数）"
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.size=Pt(12)
            ch.has_legend=True; ch.legend.position=XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout=False
            for sr in ch.plots[0].series: pass
            pl=ch.plots[0]; pl.gap_width=60; pl.overlap=100
            pl.has_data_labels=True; pl.data_labels.font.size=Pt(9); pl.data_labels.show_value=True
            try:
                sers=ch.plots[0].series
                sers[0].format.fill.solid(); sers[0].format.fill.fore_color.rgb=C_EXT   # 外部蓝
                sers[1].format.fill.solid(); sers[1].format.fill.fore_color.rgb=C_IN    # 内部红
                ch.category_axis.tick_labels.font.size=Pt(10); ch.value_axis.tick_labels.font.size=Pt(9)
            except Exception: pass
            det=[]
            for k in cats:
                v1,v2=ci.get(k,0),ce.get(k,0)
                det.append(f"{k}{v1+v2}(内{v1}/外{v2})" if v1 else f"{k}{v2}(外)")
            tb(s,Inches(8.0),Inches(3.5),Inches(4.9),Inches(1.5),"类型明细："+"；".join(det),10,False,RGBColor(0x70,0x70,0x70))
            tb(s,Inches(8.0),Inches(5.05),Inches(4.9),Inches(1.8),
               ("流失者构成：仅8月首现 %d 人｜1-7月已有记录 %d 人｜其中修过Bug %d 人｜纯评论型 %d 人"%(o['q1'][0],o['q1'][1],o['fixers'],o['cats'].get('评论型',0)) if lv=="D2" else
                "流失者构成：仅8月首现 %d 人｜1-7月已有记录 %d 人｜其中修过Bug %d 人｜纯评论型 %d 人%s"%(o['q1'][0],o['q1'][1],o['fixers'],TM[c]['pure_n'],('（其中社区任务型评论 %d 人，占%d%%，高频词「%s」）'%(TM[c]['task_n'],TM[c]['pct'],'」「'.join(TM[c]['words'])) if TM[c]['words'] else '（其中社区任务型评论 0 人，多为真实诉求而非任务党）'))),
               12,True,RGBColor(0x40,0x40,0x40))
        else:
            tb(s,Inches(4.2),Inches(4.2),Inches(5.0),Inches(1.0),f"9月无{lv}层流失",20,True,RGBColor(0x1E,0x5B,0x2A),PP_ALIGN.CENTER)
        extl=o["lost_ext_list"]
        names="、".join(extl[:28])+(("…等"+str(len(extl))+"人(全量见附录)") if len(extl)>28 else "")
        if extl: tb(s,Inches(0.4),Inches(6.95),Inches(12.5),Inches(0.5),f"{lv}外部流失名单："+names,10,False,RGBColor(0x99,0x99,0x99))

# ---- 总览 & 附录（沿用V10逻辑） ----
prs2=prs  # 简化: 总览与附录数据未变，复用同workbook
s=prs2.slides.add_slide(blank)
tb(s,Inches(0.35),Inches(0.25),Inches(12.6),Inches(0.7),"9社区总览 · 当月活跃池逐人点名（datastat 2026-10-21 快照）",21,True,RGBColor(0x1F,0x38,0x64),PP_ALIGN.CENTER)
hdr=["社区","D0·8月","D0·9月","D1池·8月","D1池·9月","D1流失(内/外)","D1新增(内/外)","D1留存率","D2池·8月","D2池·9月","D2流失(内/外)","D2新增(内/外)","D2留存率"]
data=[hdr]
for c in COMMS:
    v=X[c]["D1"]; w=X[c]["D2"]; d=D0.get(c,[0]*9)
    data.append([c,d[7],d[8],v["p8"],v["p9"],f'{v["lost_n"]}({v["lost_in"]}/{v["lost_ext"]})',f'{v["new_n"]}({v["new_in"]}/{v["new_ext"]})',pct(v["stay"],v["p8"]),
                 w["p8"],w["p9"],f'{w["lost_n"]}({w["lost_in"]}/{w["lost_ext"]})',f'{w["new_n"]}({w["new_in"]}/{w["new_ext"]})',pct(w["stay"],w["p8"])])
tot=["合计"]
for i in range(1,13):
    vals=[r[i] for r in data[1:] if isinstance(r[i],int)]
    tot.append(sum(vals) if vals else "—")
data.append(tot)
tbl=s.shapes.add_table(len(data),len(hdr),Inches(0.2),Inches(1.05),Inches(12.95),Inches(4.1)).table
for ri,rw in enumerate(data):
    for cv,v in enumerate(rw):
        cell=tbl.cell(ri,cv); cell.text=str(v)
        for p in cell.text_frame.paragraphs:
            for r0 in p.runs:
                r0.font.size=Pt(10); r0.font.bold=(ri==0 or ri==len(data)-1)
                if ri==0: r0.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
                elif cv in (5,10): r0.font.color.rgb=RGBColor(0x9C,0x00,0x06)
                elif cv in (6,11): r0.font.color.rgb=RGBColor(0x1F,0x4E,0x9C)
        if ri==0: cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor(0x1F,0x38,0x64)
tb(s,Inches(0.2),Inches(5.5),Inches(12.95),Inches(1.6),"口径：datastat『当期活跃』逐人点名；D1池=D1−D2互斥；流失=8月在该池、9月整个D1池(含D2)无记录；新增=9月入该池、8月完全不在池；柱状图为双色堆叠（■内部红/■外部蓝，段上数字=人数），『纯提PR未合入』柱仅D1页显示（提交PR未被合入、且无评论/Issue活动）。",10.5,False,RGBColor(0x70,0x70,0x70))
def list_slide(title,cols,rows_t):
    s=prs.slides.add_slide(blank)
    tb(s,Inches(0.4),Inches(0.15),Inches(12.5),Inches(0.5),title,19,True,RGBColor(0x1F,0x38,0x64),PP_ALIGN.CENTER)
    tbl=s.shapes.add_table(len(rows_t)+1,len(cols),Inches(0.25),Inches(0.95),Inches(12.85),Inches(6.3)).table
    tbl.columns[0].width=Inches(1.2)
    for cv in range(1,len(cols)): tbl.columns[cv].width=Inches(11.65/(len(cols)-1))
    for cv,h in enumerate(cols): tbl.cell(0,cv).text=h
    for ri,rw in enumerate(rows_t,1):
        for cv,v in enumerate(rw): tbl.cell(ri,cv).text=str(v)
    for ri in range(len(rows_t)+1):
        for cv in range(len(cols)):
            cell=tbl.cell(ri,cv)
            for p in cell.text_frame.paragraphs:
                for r0 in p.runs:
                    r0.font.size=Pt(10 if ri==0 else 7); r0.font.bold=(ri==0)
                    if ri==0: r0.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
        if ri==0:
            for cv in range(len(cols)):
                c0=tbl.cell(0,cv); c0.fill.solid(); c0.fill.fore_color.rgb=RGBColor(0x1F,0x38,0x64)
rows_t=[[c,X[c]["D1"]["lost_n"],f'内{X[c]["D1"]["lost_in"]}/外{X[c]["D1"]["lost_ext"]}',"、".join(sorted(set(X[c]["D1"]["lost"]))) or "—"] for c in COMMS]
list_slide("附录 · 各社区 D1 流失开发者全名单（逐人点名·含内部）",["社区","流失人数","内/外","名单"],rows_t)
rows_t=[[c,X[c]["D2"]["lost_n"],f'内{X[c]["D2"]["lost_in"]}/外{X[c]["D2"]["lost_ext"]}',"、".join(sorted(X[c]["D2"]["lost"])) or "—"] for c in COMMS]
list_slide("附录 · 各社区 D2 流失开发者全名单",["社区","流失人数","内/外","名单"],rows_t)
rows_t=[[c,X[c]["D1"]["new_n"],f'内{X[c]["D1"]["new_in"]}/外{X[c]["D1"]["new_ext"]}',"、".join(t[0] for t in X[c]["D1"]["new_list"]) or "—"] for c in COMMS]
list_slide("附录 · 各社区 9月新增 D1 名单",["社区","新增人数","内/外","名单"],rows_t)
rows_t=[[c,X[c]["D2"]["new_n"],f'内{X[c]["D2"]["new_in"]}/外{X[c]["D2"]["new_ext"]}',"、".join(t[0] for t in X[c]["D2"]["new_list"]) or "—"] for c in COMMS]
list_slide("附录 · 各社区 9月新增 D2 名单",["社区","新增人数","内/外","名单"],rows_t)
prs.save(OUT)
print("saved",OUT,os.path.getsize(OUT))
# 控制台校验：vllm D1 纯提PR数
ci,ce=GRP["vllm"]["D1"]; print("vllm D1 纯提PR未合入: 内",ci.get("纯提PR未合入",0),"外",ce.get("纯提PR未合入",0))
ci,ce=GRP["MindSpeed"]["D1"]; print("MindSpeed D1 纯提PR未合入: 内",ci.get("纯提PR未合入",0),"外",ce.get("纯提PR未合入",0))
