# -*- coding: utf-8 -*-
# x41: V9 终版 —— active池逐人点名口径 (流失/新增/构成全部可验证)
import json, os, csv
from collections import Counter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

B=__import__("os").path.join(__import__("os").path.dirname(__import__("os").path.abspath(__file__)),"data")+__import__("os").sep
DEL="/Users/tian/Desktop/开源/昇腾社区26年1月-9月贡献者分析/最新交付/"
OUT=DEL+"9社区流失开发者分析_V9_平台当前值.pptx"
COMMS=["vllm","sgl","VeRL","triton","MindIE","MindSDK","MindStudio","MindCluster","MindSpeed"]
X=json.load(open(B+"x40_data.json")); CATS=[c for c in X.pop("CATS")]
AL=json.load(open(B+"active_latest.json"))
rows=list(csv.reader(open(B+"gs_d012.csv",encoding='utf-8')))
D0={}
for r in rows[2:11]:
    nm=(r[0] or '').strip()
    if nm: D0[{'vLLM':'vllm','verl':'VeRL'}.get(nm,nm)]=[int(r[12+i] or 0) for i in range(9)]

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
def tb(s,l,tp,w,h,text,size=14,bold=False,color=RGBColor(0x33,0x33,0x33),align=PP_ALIGN.LEFT,paras=None):
    box=s.shapes.add_textbox(l,tp,w,h); tf=box.text_frame; tf.word_wrap=True
    for i,seg in enumerate(paras if paras else [text]):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        r=p.add_run(); r.text=seg; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box
def pct(a,b): return f"{a*100.0/b:.1f}%" if b else "—"

for c in COMMS:
    for lv in ("D1","D2"):
        o=X[c][lv]
        s=prs.slides.add_slide(blank)
        tb(s,Inches(0.4),Inches(0.2),Inches(12.5),Inches(0.6),f"{c} — {lv} 活跃开发者流失分析（8月→9月）",28,True,RGBColor(0x1F,0x38,0x64))
        ex8=o["p8"]-o["p8_in"]; ex9=o["p9"]-o["p9_in"]; dn=o["p9"]-o["p8"]
        red=[f"{lv} 活跃开发者总数：8月 {o['p8']} 人（内部 {o['p8_in']} 人、外部 {ex8} 人）→ 9月 {o['p9']} 人（内部 {o['p9_in']} 人、外部 {ex9} 人）",
             f"{lv} 活跃开发者9月流失：-{o['lost_n']} 人（内部 -{o['lost_in']} 人、外部 -{o['lost_ext']} 人）　|　净变化 {dn:+d} 人（{'升入D2 '+str(o.get('up',0))+'人、自D2降级流入 '+str(o.get('down',0))+'人' if lv=='D1' else '自D1升入 '+str(o.get('in_up',0))+'人、降级流出至D1 '+str(o.get('out_down',0))+'人'}，层间净{'流出' if (dn-(o['new_n']-o['lost_n']))<0 else '流入'} {abs(dn-(o['new_n']-o['lost_n']))} 人）",
             f"{lv} 活跃开发者9月新增：{o['new_n']} 人（内部 {o['new_in']} 人、外部 {o['new_ext']} 人）"]
        tb(s,Inches(0.4),Inches(0.88),Inches(12.6),Inches(2.2),None,26,True,RGBColor(0xC0,0x00,0x00),paras=red)
        tb(s,Inches(0.4),Inches(3.05),Inches(12.6),Inches(0.4),
           f"口径：当月活跃池逐人点名（D1＝PR/Issue/评论任一，D2＝有合入PR）。留存 {o['stay']} 人（{pct(o['stay'],o['p8'])}）；账目：8月池 {o['p8']} ＝ 留存{o['stay']}＋流失{o['lost_n']}＋{'升入D2 '+str(o.get('up',0)) if lv=='D1' else '降级流出 '+str(o.get('out_down',0))}；9月池 {o['p9']} ＝ 留存{o['stay']}＋新增{o['new_n']}＋{'自D2降级 '+str(o.get('down',0)) if lv=='D1' else '自D1升入 '+str(o.get('in_up',0))}",
           11,False,RGBColor(0x88,0x88,0x88))
        ordered=[(k,o["cats"][k]) for k in CATS if o["cats"].get(k,0)>0]
        if ordered:
            cd=CategoryChartData(); cd.categories=[k for k,_ in reversed(ordered)]
            cd.add_series("8月贡献活动",[v for _,v in reversed(ordered)])
            gf=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(0.4),Inches(3.5),Inches(7.4),Inches(3.35),cd)
            ch=gf.chart; ch.has_title=True
            ch.chart_title.text_frame.text=f"9月流失开发者 · 8月贡献活动分布（流失 {o['lost_n']} 人全部点名）"
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.size=Pt(13); ch.has_legend=False
            pl=ch.plots[0]; pl.has_data_labels=True; pl.data_labels.font.size=Pt(10); pl.data_labels.font.bold=True
            pl.data_labels.show_category_name=False; pl.data_labels.show_value=True; pl.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
            try:
                ch.category_axis.tick_labels.font.size=Pt(10); ch.value_axis.tick_labels.font.size=Pt(9); ch.category_axis.has_major_gridlines=False
            except Exception: pass
            tb(s,Inches(8.0),Inches(3.5),Inches(4.9),Inches(1.4),"类型明细："+"；".join(f"{k}{v}" for k,v in ordered)+f"（合计{sum(v for _,v in ordered)}人次，一人可命中多类）",11,False,RGBColor(0x70,0x70,0x70))
            tb(s,Inches(8.0),Inches(4.9),Inches(4.9),Inches(1.9),
               f"流失者构成：仅8月首现 {o['q1'][0]} 人｜1-7月已有记录 {o['q1'][1]} 人｜其中修过Bug {o['fixers']} 人｜纯评论型 {o['cats'].get('评论型',0)} 人",
               12,True,RGBColor(0x40,0x40,0x40))
        else:
            tb(s,Inches(4.2),Inches(4.2),Inches(5.0),Inches(1.0),f"9月无{lv}层流失",20,True,RGBColor(0x1E,0x5B,0x2A),PP_ALIGN.CENTER)
        extl=o["lost_ext_list"]
        names="、".join(extl[:28])+(("…等"+str(len(extl))+"人(全量见附录)") if len(extl)>28 else "")
        if extl: tb(s,Inches(0.4),Inches(6.95),Inches(12.5),Inches(0.5),f"{lv}外部流失名单："+names,10,False,RGBColor(0x99,0x99,0x99))

# ---- 总览：点名账 ----
s=prs.slides.add_slide(blank)
tb(s,Inches(0.35),Inches(0.25),Inches(12.6),Inches(0.7),"9社区总览 · 当月活跃池逐人点名（datastat 10/21，与报表数一致）",21,True,RGBColor(0x1F,0x38,0x64),PP_ALIGN.CENTER)
hdr=["社区","D0·8月","D0·9月","D1池·8月","D1池·9月","D1流失(内/外)","D1新增(内/外)","D1留存率","D2池·8月","D2池·9月","D2流失(内/外)","D2新增(内/外)","D2留存率"]
data=[hdr]
for c in COMMS:
    v=X[c]["D1"]; w=X[c]["D2"]; d=D0.get(c,[0]*9)
    data.append([c,d[7],d[8],v["p8"],v["p9"],f'{v["lost_n"]}({v["lost_in"]}/{v["lost_ext"]})',f'{v["new_n"]}({v["new_in"]}/{v["new_ext"]})',pct(v["stay"],v["p8"]),
                 w["p8"],w["p9"],f'{w["lost_n"]}({w["lost_in"]}/{w["lost_ext"]})',f'{w["new_n"]}({w["new_in"]}/{w["new_ext"]})',pct(w["stay"],w["p8"])])
tot=["合计",sum(r[i] for r in data[1:] for i in [1] if isinstance(r[i],int))]
tot=[ "合计"]
for i in range(1,13):
    vals=[r[i] for r in data[1:] if isinstance(r[i],int)]
    tot.append(sum(vals) if vals else "—")
data.append(tot)
tbl=s.shapes.add_table(len(data),len(hdr),Inches(0.2),Inches(1.05),Inches(12.95),Inches(4.1)).table
for ri,rw in enumerate(data):
    for ci,v in enumerate(rw):
        cell=tbl.cell(ri,ci); cell.text=str(v)
        for p in cell.text_frame.paragraphs:
            for r0 in p.runs:
                r0.font.size=Pt(10); r0.font.bold=(ri==0 or ri==len(data)-1)
                if ri==0: r0.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
                elif ci in (5,10) : r0.font.color.rgb=RGBColor(0x9C,0x00,0x06)
                elif ci in (6,11): r0.font.color.rgb=RGBColor(0x1F,0x4E,0x9C)
        if ri==0: cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor(0x1F,0x38,0x64)
tb(s,Inches(0.2),Inches(5.4),Inches(12.95),Inches(1.9),
 "口径声明（本版全部为逐人点名实数，非推导）：当月活跃池=level/detail『当期活跃』口径（与报表活跃数逐一对上，如 vllm D1 8月754→池754人）；D1池为互斥值(=D1−D2)。流失=8月在该池、9月在整个D1池(含D2)无任何活动记录；新增=9月出现在该池、8月完全不在池；两者用账号集合差点名，流失名单见附录。留存率=8月池成员9月仍在同池占比。注：单月活跃池天然高换手（多数人非月月活跃），留存率低≠关系破裂。",
 10.5,False,RGBColor(0x70,0x70,0x70))

# ---- 附录 ----
def list_slide(title,cols,rows_t,sub=None):
    s=prs.slides.add_slide(blank)
    tb(s,Inches(0.4),Inches(0.15),Inches(12.5),Inches(0.5),title,19,True,RGBColor(0x1F,0x38,0x64),PP_ALIGN.CENTER)
    ty=0.95
    if sub:
        tb(s,Inches(0.3),Inches(0.6),Inches(12.75),Inches(0.6),sub,10,False,RGBColor(0x88,0x60,0x00),PP_ALIGN.CENTER); ty=1.25
    tbl=s.shapes.add_table(len(rows_t)+1,len(cols),Inches(0.25),Inches(ty),Inches(12.85),Inches(7.25-ty)).table
    tbl.columns[0].width=Inches(1.2)
    for ci in range(1,len(cols)): tbl.columns[ci].width=Inches(11.65/(len(cols)-1))
    for ci,h in enumerate(cols): tbl.cell(0,ci).text=h
    for ri,rw in enumerate(rows_t,1):
        for ci,v in enumerate(rw): tbl.cell(ri,ci).text=str(v)
    for ri in range(len(rows_t)+1):
        for ci in range(len(cols)):
            cell=tbl.cell(ri,ci)
            for p in cell.text_frame.paragraphs:
                for r0 in p.runs:
                    r0.font.size=Pt(10 if ri==0 else 7); r0.font.bold=(ri==0)
                    if ri==0: r0.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
        if ri==0:
            for ci in range(len(cols)):
                c0=tbl.cell(0,ci); c0.fill.solid(); c0.fill.fore_color.rgb=RGBColor(0x1F,0x38,0x64)
    return s
rows_t=[[c,X[c]["D1"]["lost_n"],f'内{X[c]["D1"]["lost_in"]}/外{X[c]["D1"]["lost_ext"]}',"、".join(sorted(set(X[c]["D1"]["lost"]))) or "—"] for c in COMMS]
list_slide("附录 · 各社区 D1 流失开发者全名单（逐人点名·含内部）",["社区","流失人数","内/外","名单"],rows_t,
 sub="名单=8月在D1池、9月整个池无记录的账号（小写归一）；纯评论型流失者也在其中——此前『无法点名的197人』实际全部在册。")
rows_t=[[c,X[c]["D2"]["lost_n"],f'内{X[c]["D2"]["lost_in"]}/外{X[c]["D2"]["lost_ext"]}',"、".join(sorted(X[c]["D2"]["lost"])) or "—"] for c in COMMS]
list_slide("附录 · 各社区 D2 流失开发者全名单",["社区","流失人数","内/外","名单"],rows_t)
rows_t=[[c,X[c]["D1"]["new_n"],f'内{X[c]["D1"]["new_in"]}/外{X[c]["D1"]["new_ext"]}',"、".join(t[0] for t in X[c]["D1"]["new_list"]) or "—"] for c in COMMS]
list_slide("附录 · 各社区 9月新增 D1 名单（8月完全不在池）",["社区","新增人数","内/外","名单"],rows_t)
rows_t=[[c,X[c]["D2"]["new_n"],f'内{X[c]["D2"]["new_in"]}/外{X[c]["D2"]["new_ext"]}',"、".join(t[0] for t in X[c]["D2"]["new_list"]) or "—"] for c in COMMS]
list_slide("附录 · 各社区 9月新增 D2 名单（8月完全不在池）",["社区","新增人数","内/外","名单"],rows_t)
prs.save(OUT)
print("saved",OUT,os.path.getsize(OUT))
